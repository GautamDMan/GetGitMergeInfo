#!/usr/bin/env python3
"""
List ALL merges on a branch of a local git repository and write them to Excel.

Reads history read-only (no checkout / reset), following the branch's own
first-parent line, so every row is a merge that actually landed on that branch
and the "files" of a merge are exactly what it brought in (diff vs 1st parent).

Workbook sheets (output.xlsx):
  1. All Merges     - every merge on the branch, newest first
  2. Last Merge     - (only if input.csv is given) latest merge per input file
  3. File History   - (only if input.csv is given) every merge that touched
                      each input file

Usage:
    python get_all_merges_excel.py --creds git_login.json --output merges.xlsx
    python get_all_merges_excel.py --creds git_login.json --input input.csv --output merges.xlsx

git_login.json:
    {
      "repo_path": "/path/to/the/repo",
      "branch": "main"          // optional, default HEAD
    }

input.csv (optional):
    file_name
    config.py
    utils.py

Optional one-time speedup for very large repos:
    git commit-graph write --reachable --changed-paths
"""

import argparse
import csv
import json
import os
import re
import subprocess
import sys
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SEP = "\x01"
MARK = "\x02"
MAX_FILES_LISTED = 200          # files shown per merge in the "files" cell
FONT = "Arial"

PR_MERGE_RE = re.compile(r"Merge pull request #\d+ from ([^\s/]+)/(\S+)")
BRANCH_MERGE_RE = re.compile(r"Merge branch '([^']+)'")
ILLEGAL_XLSX_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


# --------------------------------------------------------------------------- #
# git helpers
# --------------------------------------------------------------------------- #
def git_cmd(repo_path, args):
    return ["git", "-C", repo_path, "-c", "core.quotepath=false"] + args


def run_git(repo_path, args, stdin_text=None, check=True):
    r = subprocess.run(
        git_cmd(repo_path, args), input=stdin_text, capture_output=True,
        text=True, encoding="utf-8", errors="replace",
    )
    if check and r.returncode != 0:
        raise RuntimeError(f"git {' '.join(args)} failed: {r.stderr.strip()}")
    return r.stdout


def parse_subject(subject):
    """Return (pr_owner_or_None, merged_branch_or_'')."""
    m = PR_MERGE_RE.search(subject)
    if m:
        return m.group(1), m.group(2)
    m = BRANCH_MERGE_RE.search(subject)
    if m:
        return None, m.group(1)
    return None, ""


def load_credentials(creds_path):
    with open(creds_path, "r", encoding="utf-8") as f:
        creds = json.load(f)
    repo_path = creds.get("repo_path")
    if not repo_path:
        raise ValueError("git_login.json must contain a 'repo_path' field")
    if not os.path.isdir(repo_path):
        raise ValueError(f"'{repo_path}' is not a directory")
    out = run_git(repo_path, ["rev-parse", "--is-inside-work-tree"], check=False).strip()
    if out != "true":
        raise ValueError(f"'{repo_path}' does not look like a git repository")
    return repo_path, creds.get("branch")


def load_file_names(input_csv):
    names = []
    with open(input_csv, "r", newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        if "file_name" not in (reader.fieldnames or []):
            raise ValueError(f"input.csv must have a 'file_name' column. Found: {reader.fieldnames}")
        for row in reader:
            n = (row["file_name"] or "").strip()
            if n:
                names.append(n)
    return list(dict.fromkeys(names))


def build_name_index(repo_path, ref):
    out = run_git(repo_path, ["ls-tree", "-r", "--name-only", "-z", ref])
    idx = defaultdict(list)
    for p in out.split("\0"):
        if p:
            idx[p.rsplit("/", 1)[-1]].append(p)
    return idx


# --------------------------------------------------------------------------- #
# one pass over the branch's merges
# --------------------------------------------------------------------------- #
def scan_merges(repo_path, ref, targets):
    """All merges on `ref` (first-parent line), newest first.

    Each merge: hash, parents, author, subject, committed_at, nfiles, files,
    hits (subset of `targets` it touched)."""
    fmt = f"{MARK}%H{SEP}%P{SEP}%an{SEP}%s{SEP}%cI"
    proc = subprocess.Popen(
        git_cmd(repo_path, ["log", ref, "--first-parent", "--merges", "-m",
                            "--name-only", f"--pretty=format:{fmt}"]),
        stdout=subprocess.PIPE, stderr=subprocess.PIPE,
        text=True, encoding="utf-8", errors="replace",
    )
    merges, cur = [], None
    for raw in proc.stdout:
        line = raw.rstrip("\n").rstrip("\r")
        if line.startswith(MARK):
            parts = line[1:].split(SEP)
            if len(parts) != 5:
                cur = None
                continue
            if cur is not None and cur["hash"] == parts[0]:
                continue  # same merge repeated by git - keep accumulating
            cur = {"hash": parts[0], "parents": parts[1], "author": parts[2],
                   "subject": parts[3], "committed_at": parts[4],
                   "nfiles": 0, "files": [], "hits": set()}
            merges.append(cur)
        elif line and cur is not None:
            cur["nfiles"] += 1
            if len(cur["files"]) < MAX_FILES_LISTED:
                cur["files"].append(line)
            if line in targets:
                cur["hits"].add(line)
    err = proc.stderr.read()
    if proc.wait() != 0:
        raise RuntimeError(f"git log failed: {err.strip()}")
    return merges


def batch_pushed_by(repo_path, merges):
    """{tip_hash: (name, email)} for the 2nd parent of every merge, one git call."""
    targets = set()
    for m in merges:
        ph = m["parents"].split()
        targets.add(ph[1] if len(ph) >= 2 else m["hash"])
    if not targets:
        return {}
    out = run_git(
        repo_path,
        ["log", "--no-walk=unsorted", "--stdin", f"--pretty=format:%H{SEP}%an{SEP}%ae"],
        stdin_text="\n".join(sorted(targets)) + "\n", check=False,
    )
    res = {}
    for line in out.splitlines():
        p = line.split(SEP)
        if len(p) == 3:
            res[p[0]] = (p[1], p[2])
    return res


def enrich(merges, pushed):
    for m in merges:
        pr_owner, merged_branch = parse_subject(m["subject"])
        ph = m["parents"].split()
        tip = ph[1] if len(ph) >= 2 else m["hash"]
        m["merge_owner"] = pr_owner or m["author"]
        m["merged_branch"] = merged_branch
        m["pushed_by"], m["pushed_by_email"] = pushed.get(tip, ("", ""))


# --------------------------------------------------------------------------- #
# Excel output
# --------------------------------------------------------------------------- #
def to_dt(iso):
    try:
        return datetime.fromisoformat(iso).replace(tzinfo=None)  # Excel has no tz support
    except ValueError:
        return iso


def write_sheet(ws, headers, rows, widths, wrap_cols=()):
    head_font = Font(name=FONT, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", start_color="305496")
    body_font = Font(name=FONT)
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.font = body_font
            if isinstance(c.value, str):
                c.value = ILLEGAL_XLSX_RE.sub("", c.value)
                if c.value.startswith("="):      # never let text become a formula
                    c.data_type = "s"
            if isinstance(c.value, datetime):
                c.number_format = "yyyy-mm-dd hh:mm:ss"
            if c.column in wrap_cols:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                c.alignment = Alignment(vertical="top")
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


def merge_row(m):
    return [to_dt(m["committed_at"]), m["merge_owner"], m["pushed_by"], m["pushed_by_email"],
            m["subject"], m["merged_branch"], m["hash"]]


def build_workbook(merges, ref, file_results):
    wb = Workbook()

    # Sheet 1: every merge on the branch
    ws = wb.active
    ws.title = "All Merges"
    rows = []
    for i, m in enumerate(merges, start=1):
        files_text = "\n".join(m["files"])
        if m["nfiles"] > len(m["files"]):
            files_text += f"\n... (+{m['nfiles'] - len(m['files'])} more)"
        rows.append([i] + merge_row(m) + [m["nfiles"], files_text])
    write_sheet(
        ws,
        ["#", "merged_at", "merge_owner", "pushed_by", "pushed_by_email", "merge_message",
         "merged_branch", "merge_commit", "files_changed", "files"],
        rows, [6, 20, 18, 18, 28, 55, 28, 42, 14, 60], wrap_cols=(10,),
    )
    ws.append([])
    ws.append([f"Branch/ref: {ref}  |  merges: {len(merges)}  |  "
               "first-parent history; files = changed vs first parent"])
    ws.cell(row=ws.max_row, column=1).font = Font(name=FONT, italic=True, color="595959")

    if file_results is None:
        return wb

    # Sheet 2: last merge per input file
    ws2 = wb.create_sheet("Last Merge")
    rows2 = []
    for fr in file_results:
        m = fr["merges"][0] if fr["merges"] else None
        if m:
            rows2.append([fr["file_name"], fr["path"]] + merge_row(m) + [len(fr["merges"])])
        else:
            rows2.append([fr["file_name"], fr["path"]] + [""] * 7 + [0])
            rows2[-1][6] = fr["error"]
    write_sheet(
        ws2,
        ["file_name", "resolved_path", "merged_at", "merge_owner", "pushed_by",
         "pushed_by_email", "merge_message", "merged_branch", "merge_commit", "total_merges"],
        rows2, [24, 40, 20, 18, 18, 28, 55, 28, 42, 14],
    )

    # Sheet 3: every merge that touched each input file
    ws3 = wb.create_sheet("File History")
    rows3 = []
    for fr in file_results:
        for m in fr["merges"]:
            rows3.append([fr["file_name"], fr["path"]] + merge_row(m))
    write_sheet(
        ws3,
        ["file_name", "resolved_path", "merged_at", "merge_owner", "pushed_by",
         "pushed_by_email", "merge_message", "merged_branch", "merge_commit"],
        rows3, [24, 40, 20, 18, 18, 28, 55, 28, 42],
    )
    return wb


# --------------------------------------------------------------------------- #
# main
# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser(description="Export all merges of a branch to Excel")
    ap.add_argument("--creds", default="git_login.json", help="Path to git_login.json")
    ap.add_argument("--input", default=None, help="Optional input.csv with a file_name column")
    ap.add_argument("--output", default="merges.xlsx", help="Output .xlsx path")
    args = ap.parse_args()

    try:
        repo_path, branch = load_credentials(args.creds)
    except Exception as e:
        print(f"Error loading credentials: {e}", file=sys.stderr)
        sys.exit(1)
    ref = branch or "HEAD"

    file_names, path_for, errors = [], {}, {}
    if args.input:
        try:
            file_names = load_file_names(args.input)
            idx = build_name_index(repo_path, ref)
        except Exception as e:
            print(f"Error preparing input files: {e}", file=sys.stderr)
            sys.exit(1)
        for n in file_names:
            hits = idx.get(n, [])
            if len(hits) == 1:
                path_for[n] = hits[0]
            elif not hits:
                errors[n] = f"ERROR: file '{n}' not found in repo"
            else:
                errors[n] = f"ERROR: ambiguous: {len(hits)} files named '{n}': {sorted(hits)}"

    print(f"Scanning merges on '{ref}' in {repo_path} (read-only)...")
    try:
        merges = scan_merges(repo_path, ref, set(path_for.values()))
        enrich(merges, batch_pushed_by(repo_path, merges))
    except Exception as e:
        print(f"Error reading history: {e}", file=sys.stderr)
        sys.exit(1)

    file_results = None
    if args.input:
        file_results = []
        for n in file_names:
            p = path_for.get(n, "")
            hit_merges = [m for m in merges if p and p in m["hits"]]
            err = errors.get(n) or ("" if hit_merges else "ERROR: no merge commit found that touched this file")
            file_results.append({"file_name": n, "path": p, "merges": hit_merges, "error": err})

    wb = build_workbook(merges, ref, file_results)
    wb.save(args.output)
    print(f"Done. {len(merges)} merges written to {args.output}")


if __name__ == "__main__":
    main()
